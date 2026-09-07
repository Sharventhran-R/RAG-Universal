"""Generate the demo corpus: one pdf, one xlsx, one docx.

Small, deterministic, and content-matched to the smoke test's three questions
(one pure-semantic, one metadata-filtered, one delete). Called by
``python -m app.cli seed`` when the files are missing, and by the smoke test.
"""

from __future__ import annotations

from pathlib import Path

FIXTURES_DIR = Path(__file__).resolve().parent


def build(dest: Path | None = None) -> dict[str, Path]:
    dest = dest or FIXTURES_DIR
    dest.mkdir(parents=True, exist_ok=True)
    return {
        "pdf": _handbook_pdf(dest / "handbook.pdf"),
        "xlsx": _sales_xlsx(dest / "sales.xlsx"),
        "docx": _notes_docx(dest / "notes.docx"),
    }


def _handbook_pdf(path: Path) -> Path:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 80), "Company Handbook", fontsize=22)
    page.insert_text((72, 120), "History", fontsize=15)
    page.insert_text(
        (72, 150),
        "Northwind Systems was founded in 1998 in the city of Rotterdam.",
        fontsize=11,
    )
    page.insert_text((72, 175), "It has operated continuously ever since.", fontsize=11)
    page.insert_text((72, 215), "Finance", fontsize=15)
    page.insert_text(
        (72, 245),
        "Annual revenue reached 42 million euros in fiscal year 2025.",
        fontsize=11,
    )
    doc.save(str(path))
    doc.close()
    return path


def _sales_xlsx(path: Path) -> Path:
    import pandas as pd

    df = pd.DataFrame(
        {
            "Region": ["North", "South", "East", "West"],
            "Revenue_MEUR": [12.0, 9.5, 11.0, 9.5],
            "Deals": [140, 95, 120, 88],
        }
    )
    with pd.ExcelWriter(path) as writer:
        df.to_excel(writer, sheet_name="FY2025", index=False)
    return path


def _notes_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Onboarding Notes", level=1)
    doc.add_paragraph("New hires should complete the security training in week one.")
    doc.add_heading("Facilities", level=2)
    doc.add_paragraph("The Rotterdam office is open from 08:00 to 19:00 on weekdays.")
    doc.save(str(path))
    return path


if __name__ == "__main__":  # pragma: no cover
    for kind, p in build().items():
        print(f"{kind:5} -> {p}")
