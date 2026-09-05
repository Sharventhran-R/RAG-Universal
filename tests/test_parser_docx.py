from __future__ import annotations

import io
from pathlib import Path

import pytest

pytest.importorskip("docx")
from docx import Document

from app.parsers.base import EMPTY_NO_TEXT, ParseInput
from app.parsers.docx import DocxParser


def _src(data: bytes, document_id: str = "doc1") -> ParseInput:
    return ParseInput(
        document_id=document_id,
        filename="review.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_sha256="sha",
        extractor_version=DocxParser.version,
        open_stream=lambda: io.BytesIO(data),
        local_path=lambda: Path("unused"),
    )


def _doc_bytes(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample() -> bytes:
    def build(doc):
        doc.add_heading("Quarterly Review", level=1)
        doc.add_paragraph("Revenue grew to 5,000 in Q3.")
        doc.add_heading("Details", level=2)
        doc.add_paragraph("First item", style="List Bullet")
        table = doc.add_table(rows=3, cols=2)
        cells = {(0, 0): "Region", (0, 1): "Amount",
                 (1, 0): "North", (1, 1): "500",
                 (2, 0): "South", (2, 1): "734"}
        for (r, c), v in cells.items():
            table.cell(r, c).text = v

    return _doc_bytes(build)


def test_identity():
    assert DocxParser.file_type == "docx"


def test_clean_document_has_no_flags(sample):
    res = DocxParser().parse(_src(sample))
    assert res.extraction_flags == []
    assert res.status_hint is None


def test_block_sequence_and_section_path(sample):
    blocks = DocxParser().parse(_src(sample)).blocks
    types = [b.type for b in blocks]
    assert types[0] == "heading" and blocks[0].content == "Quarterly Review"
    assert blocks[0].section_path == []  # a heading's own path is its ancestors

    para = next(b for b in blocks if b.type == "paragraph")
    assert "5,000" in para.content
    assert para.section_path == ["Quarterly Review"]

    li = next(b for b in blocks if b.type == "list_item")
    assert li.section_path == ["Quarterly Review", "Details"]

    tbl = next(b for b in blocks if b.type == "table")
    assert "| Region | Amount |" in tbl.content and "734" in tbl.content
    assert tbl.section_path == ["Quarterly Review", "Details"]
    assert tbl.flag is None


def test_no_page_locator_but_char_spans_are_ordered(sample):
    blocks = DocxParser().parse(_src(sample)).blocks
    prev_end = -1
    for b in blocks:
        assert b.locator.page is None
        assert b.locator.char_start is not None and b.locator.char_end is not None
        assert b.locator.char_start >= prev_end
        prev_end = b.locator.char_end


def test_ids_are_deterministic(sample):
    a = [b.id for b in DocxParser().parse(_src(sample)).blocks]
    b = [b.id for b in DocxParser().parse(_src(sample)).blocks]
    assert a == b


def test_layout_table_is_flattened_to_text_and_flagged():
    def build(doc):
        doc.add_paragraph("Intro")
        t = doc.add_table(rows=1, cols=1)          # 1x1 -> not a real table
        t.cell(0, 0).text = "just a boxed sentence"

    res = DocxParser().parse(_src(_doc_bytes(build)))
    assert "low_confidence_table" in res.extraction_flags
    assert not any(b.type == "table" for b in res.blocks)
    assert any(b.content == "just a boxed sentence" and b.type == "paragraph" for b in res.blocks)


def test_empty_document_signals_empty_no_text():
    res = DocxParser().parse(_src(_doc_bytes(lambda d: None)))
    assert res.blocks == []
    assert res.status_hint == EMPTY_NO_TEXT
    assert "docx_no_text" in res.extraction_flags
